#!/usr/bin/env python3
"""Parse Birthday _ Anniversary - Bubbles & Black Tie (1).xlsx into Strapi seed JS."""

from __future__ import annotations

import json
import re
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "Birthday _ Anniversary - Bubbles & Black Tie (1).xlsx"
JSON_OUT = ROOT / "tmp" / "bubbles_packages_parsed.json"
SEED_OUT = ROOT / "strapi" / "src" / "seed" / "decor-bubbles-black-tie-packages.js"

# Keep existing production slugs so refresh updates in place (no duplicates).
META: dict[str, tuple[str, str, str, str, str]] = {
    "GT 1": ("golden-touch-1", "The Golden Touch 1", "The Golden Touch", "level-1", "golden-touch"),
    "GT 2": ("golden-touch-2", "The Golden Touch 2", "The Golden Touch", "level-2", "golden-touch"),
    "GT 23 ": ("golden-touch-2-3", "The Golden Touch 2/3", "The Golden Touch", "level-2", "golden-touch"),
    "GT 3": ("golden-touch-3", "The Golden Touch 3", "The Golden Touch", "level-3", "golden-touch"),
    "SS 1 ": ("social-setup-1", "The Social Setup 1", "The Social Setup", "level-1", "social-setup"),
    "SS2 ": ("social-setup-2", "The Social Setup 2", "The Social Setup", "level-2", "social-setup"),
    "SS 23": ("social-setup-2-3", "The Social Setup 2/3", "The Social Setup", "level-2", "social-setup"),
    "SS3": ("social-setup-3", "The Social Setup 3", "The Social Setup", "level-3", "social-setup"),
    "CP": (
        "complete-package-bubbles-black-tie",
        "Complete Package",
        "Complete Package",
        "level-3",
        "complete-package",
    ),
}

DEFAULT_NITTY = (
    "Love the look but need a different color? go ahead and book! "
    "We'll contact you to finalize the perfect shade from our available decor to match your vibe."
)
DEFAULT_SETUP = "A small area for our team to work, access to water and electricity."
PRICE_COVERS = [
    {"text": "The Complete Look"},
    {"text": "Professional Setup"},
    {"text": "No hidden service fees."},
]

# Short card chips (Baby Shower / Pastel style) — never dump Excel Includes paragraphs.
SHORT_CARD_TAGS: dict[str, list[str]] = {
    "GT 1": ["5 golden floral centrepieces", "Entrance easel + board"],
    "GT 2": [
        "5 golden floral centrepieces",
        "Entrance easel + florals",
        "Basic bar setup",
        "Candles in glass",
    ],
    "GT 23 ": [
        "5 premium floral centrepieces",
        "Entrance easel + florals",
        "Basic bar setup",
        "Backdrop with balloons",
        "Candles in glass",
    ],
    "GT 3": [
        "5 premium floral centrepieces",
        "Golden frame entrance",
        "Premium bar + menu",
        "Double panel backdrop",
        "Candles in glass jars",
    ],
    "SS 1 ": [
        "5 premium floral centrepieces",
        "3 cocktail centrepieces",
        "Entrance easel + florals",
        "Basic bar setup",
        "Backdrop with balloons",
        "Tables + chairs set",
        "Fairy lights",
    ],
    "SS2 ": [
        "5 premium floral centrepieces",
        "Entrance easel + florals",
        "Basic bar setup",
        "Backdrop with balloons",
        "Tables + chairs set",
        "Fairy lights",
        "Printables / games",
        "LED / halogen lights",
    ],
    "SS 23": [
        "5 premium floral centrepieces",
        "3 cocktail centrepieces",
        "Golden frame entrance",
        "Premium bar + menu",
        "Double panel backdrop",
        "Premium furniture set",
        "Fairy lights + lanterns",
        "Themed props",
    ],
    "SS3": [
        "5 premium floral centrepieces",
        "Golden frame entrance",
        "Premium bar + menu",
        "Double panel backdrop",
        "Premium furniture set",
        "Fairy lights + lanterns",
        "Printables",
        "LED lights + props",
    ],
    "CP": [
        "5 premium floral centrepieces",
        "Golden frame entrance",
        "Premium bar + menu",
        "Double panel backdrop",
        "Furniture set",
        "Fairy lights + lanterns",
        "Printables + props",
        "Catering for 35 pax",
    ],
}

SECTION_STOPS = [
    "Includes:",
    "Includes",
    "The Nitty-Gritty:",
    "The Nitty-Gritty",
    "Nitty - Gritties",
    "The price you see covers it all:",
    "The price you see covers it all",
    "Best Placed at",
    "Best placed",
    "Why we Love it",
    "Why we love",
    "Setup Needs",
    "Setup",
    "Policies",
    "Cancellation",
    "Key Decor Elements:",
    "Key Decor Elements",
    "Reference",
]


def norm(value: object) -> str:
    if value is None:
        return ""
    text = str(value).replace("\r\n", "\n").replace("\r", "\n").strip()
    return re.sub(r"[ \t]+", " ", text)


def amt(value: object) -> float | None:
    if isinstance(value, (int, float)):
        return float(value)
    text = (
        norm(value)
        .replace(",", "")
        .replace("₹", "")
        .replace("Rs.", "")
        .replace("Rs", "")
        .replace("INR", "")
    )
    match = re.search(r"-?\d+(?:\.\d+)?", text)
    return float(match.group(0)) if match else None


def find_label(ws, labels: list[str], start: int = 1) -> int | None:
    targets = {label.lower().rstrip(":") for label in labels}
    for row in range(start, (ws.max_row or 1) + 1):
        cell = norm(ws.cell(row, 1).value).lower().rstrip(":")
        if cell in targets:
            return row
    return None


def amount_on_row(ws, row: int) -> float | None:
    for col in range(2, 8):
        value = amt(ws.cell(row, col).value)
        if value is not None:
            return value
    return None


def section_after(ws, labels: list[str], stop_labels: list[str]) -> str:
    row = find_label(ws, labels)
    if row is None:
        return ""
    stops = {label.lower().rstrip(":") for label in stop_labels}
    lines: list[str] = []
    for next_row in range(row + 1, (ws.max_row or row) + 1):
        text = norm(ws.cell(next_row, 1).value)
        if not text:
            continue
        key = text.lower().rstrip(":")
        if key in stops or any(key.startswith(stop) for stop in stops):
            break
        if key in ("description", "decription", "total price"):
            continue
        lines.append(text)
    return "\n".join(lines)


def clean_prose(text: str, fallback: str = "") -> str:
    cleaned = (text or "").strip()
    if not cleaned:
        return fallback
    lower = cleaned.lower()
    for marker in (
        " policies ",
        "policies all the material",
        "cancellation policy",
        "read our full",
        "[read our",
        "reference picture",
    ):
        idx = lower.find(marker)
        if idx > 0:
            cleaned = cleaned[:idx].strip(" .")
            break
    return cleaned[:500] if cleaned else fallback


def js_string(value: str) -> str:
    escaped = (
        value.replace("\\", "\\\\")
        .replace("'", "\\'")
        .replace("\n", "\\n")
        .replace("\r", "")
    )
    return "'%s'" % escaped


def parse_sheet(ws, sheet_name: str, meta: tuple[str, str, str, str, str]) -> dict:
    slug, title, theme_name, level, child = meta
    package_price = platform = labour = total = None

    for row in range(1, (ws.max_row or 1) + 1):
        label = norm(ws.cell(row, 1).value).lower().rstrip(":")
        if label in ("package price", "total") and label != "total price":
            # Prefer numeric "Total" before fees when present.
            if label == "total":
                value = amount_on_row(ws, row)
                if value is not None and value >= 1000:
                    package_price = value
            elif label == "package price":
                package_price = amount_on_row(ws, row)
        elif label in ("transport",) or ("transport" in label and "labour" in label):
            labour = amount_on_row(ws, row)
        elif label == "platform fee":
            platform = amount_on_row(ws, row)
        elif label.startswith("total price"):
            total = amount_on_row(ws, row)

    desc_raw = section_after(
        ws,
        ["Description", "Decription", "Description "],
        SECTION_STOPS,
    )
    if not desc_raw:
        # Some sheets put description as a bare paragraph after pricing.
        for row in range(1, min(30, (ws.max_row or 1) + 1)):
            text = norm(ws.cell(row, 1).value)
            if len(text) > 80 and not text.lower().startswith(
                ("includes", "the nitty", "centerpieces", "total")
            ):
                desc_raw = text
                break

    description = clean_prose(
        " ".join(line for line in desc_raw.split("\n") if line and not line.startswith("-")).strip(),
        title,
    )

    key_raw = section_after(ws, ["Key Decor Elements:", "Key Decor Elements"], SECTION_STOPS)
    key_decor = clean_prose(" ".join(key_raw.split("\n")).strip(), description[:180])

    inclusions = list(SHORT_CARD_TAGS.get(sheet_name, []))

    nitty = section_after(
        ws,
        ["The Nitty-Gritty:", "The Nitty-Gritty", "Nitty - Gritties"],
        SECTION_STOPS,
    )
    nitty = clean_prose(" ".join(nitty.split("\n")).strip(), DEFAULT_NITTY)

    best = section_after(ws, ["Best Placed at", "Best placed"], SECTION_STOPS)
    best = clean_prose(" ".join(best.split("\n")).strip(), "")

    why = section_after(ws, ["Why we Love it", "Why we love it", "Why we love"], SECTION_STOPS)
    why = clean_prose(" ".join(why.split("\n")).strip(), "")

    setup = section_after(ws, ["Setup Needs", "Setup"], SECTION_STOPS)
    setup = clean_prose(" ".join(setup.split("\n")).strip(), DEFAULT_SETUP)

    breakdown: list[dict[str, int]] = []
    for row in range(2, 40):
        label = norm(ws.cell(row, 1).value)
        label_key = label.lower().rstrip(":")
        if label_key in ("total", "platform fee", "transport") or label_key.startswith(
            "total price"
        ):
            break
        detail = norm(ws.cell(row, 2).value) or label
        cost = amt(ws.cell(row, 3).value)
        if (
            detail
            and cost is not None
            and cost >= 100
            and detail.lower() not in ("details", "estimated cost", "average")
        ):
            breakdown.append({"label": (label or detail)[:80], "amount": int(round(cost))})

    if platform is not None:
        breakdown.append({"label": "Platform fee", "amount": int(round(platform))})
    if labour is not None:
        breakdown.append({"label": "Labour & Transport", "amount": int(round(labour))})

    if total is not None and total > 0:
        base_price = int(round(total))
    else:
        base_price = int(round((package_price or 0) + (platform or 0) + (labour or 0)))

    if base_price <= 0:
        raise ValueError("Sheet %r has no usable total price" % sheet_name)

    return {
        "sheet": sheet_name.strip(),
        "title": title,
        "slug": slug,
        "theme_name": theme_name,
        "budget_level_slug": level,
        "base_price": base_price,
        "description": description[:1000],
        "key_decor_elements": key_decor[:180],
        "nitty_gritty": nitty[:800],
        "setup_needs": setup[:500],
        "best_placed_at": best[:500],
        "why_we_love_it": why[:500],
        "child_theme": child,
        "price_breakdown": breakdown,
        "inclustions": inclusions[:12],
    }


def render_seed(packages: list[dict]) -> str:
    chunks: list[str] = [
        "'use strict';",
        "",
        "const {",
        "  ensureBudgetLevelDocumentId,",
        "  ensureEventDocumentId,",
        "  ensureThemeDocumentId,",
        "  seedDecorPackages,",
        "} = require('./decor-proposal-shared');",
        "",
        "const EVENT_SLUG_CANDIDATES = [",
        "  'birthday-parties',",
        "  'birthday-celebrations',",
        "  'birthday',",
        "  'birthday-party',",
        "  'anniversary',",
        "  'anniversary-evenings',",
        "  'anniversaries',",
        "];",
        "",
        "const BUBBLES_BLACK_TIE_PACKAGES = [",
    ]

    for pkg in packages:
        lines = ["  {"]
        lines.append("    sheet: %s," % js_string(pkg["sheet"]))
        lines.append("    title: %s," % js_string(pkg["title"]))
        lines.append("    slug: %s," % js_string(pkg["slug"]))
        lines.append("    theme_name: %s," % js_string(pkg["theme_name"]))
        lines.append("    budget_level_slug: %s," % js_string(pkg["budget_level_slug"]))
        lines.append("    base_price: %s," % pkg["base_price"])
        lines.append("    description: %s," % js_string(pkg["description"]))
        lines.append("    key_decor_elements: %s," % js_string(pkg["key_decor_elements"]))
        lines.append("    nitty_gritty: %s," % js_string(pkg["nitty_gritty"]))
        lines.append("    setup_needs: %s," % js_string(pkg["setup_needs"]))
        lines.append("    best_placed_at: %s," % js_string(pkg["best_placed_at"]))
        lines.append("    why_we_love_it: %s," % js_string(pkg["why_we_love_it"]))
        lines.append(
            "    theme_slug_candidates: ['bubbles-and-black-tie', %s],"
            % js_string(pkg["child_theme"])
        )
        lines.append("    price_covers: [")
        for cover in PRICE_COVERS:
            lines.append("      { text: %s }," % js_string(cover["text"]))
        lines.append("    ],")
        lines.append("    price_breakdown: [")
        for item in pkg["price_breakdown"]:
            lines.append(
                "      { label: %s, amount: %s },"
                % (js_string(item["label"]), item["amount"])
            )
        lines.append("    ],")
        lines.append("    inclustions: [")
        for text in pkg["inclustions"]:
            lines.append(
                "      { text: %s, group_title: 'Includes' }," % js_string(text)
            )
        lines.append("    ],")
        lines.append("  },")
        chunks.append("\n".join(lines))

    chunks.extend(
        [
            "];",
            "",
            "async function seedDecorBubblesBlackTiePackages(strapi) {",
            "  await ensureThemeDocumentId(strapi, {",
            "    name: 'Bubbles & Black Tie',",
            "    slug: 'bubbles-and-black-tie',",
            "    display_order: 40,",
            "  });",
            "  await ensureEventDocumentId(strapi, {",
            "    name: 'Birthday Parties',",
            "    slugCandidates: EVENT_SLUG_CANDIDATES,",
            "    display_order: 10,",
            "  });",
            "  await ensureBudgetLevelDocumentId(strapi, {",
            "    name: 'Level 1',",
            "    slug: 'level-1',",
            "    display_order: 1,",
            "    description: 'Thoughtfully put together essentials for smaller guest counts.',",
            "  });",
            "  await ensureBudgetLevelDocumentId(strapi, {",
            "    name: 'Level 2',",
            "    slug: 'level-2',",
            "    display_order: 2,",
            "    description: 'A more detailed setup with a stronger visual and service experience.',",
            "  });",
            "  await ensureBudgetLevelDocumentId(strapi, {",
            "    name: 'Level 3',",
            "    slug: 'level-3',",
            "    display_order: 3,",
            "    description: 'Premium styling and service depth for larger or more layered celebrations.',",
            "  });",
            "",
            "  return seedDecorPackages(strapi, BUBBLES_BLACK_TIE_PACKAGES, {",
            "    eventSlugCandidates: EVENT_SLUG_CANDIDATES,",
            "    logPrefix: 'seed/decor-bubbles-black-tie',",
            "    refreshExistingDecor: true,",
            "  });",
            "}",
            "",
            "module.exports = {",
            "  BUBBLES_BLACK_TIE_PACKAGES,",
            "  EVENT_SLUG_CANDIDATES,",
            "  seedDecorBubblesBlackTiePackages,",
            "};",
            "",
        ]
    )
    return "\n".join(chunks)


def main() -> None:
    workbook = load_workbook(XLSX, data_only = True)
    packages: list[dict] = []
    seen_slugs: set[str] = set()

    for sheet_name, meta in META.items():
        if sheet_name not in workbook.sheetnames:
            raise SystemExit("Missing sheet: %r" % sheet_name)
        pkg = parse_sheet(workbook[sheet_name], sheet_name, meta)
        if pkg["slug"] in seen_slugs:
            raise SystemExit("Duplicate slug: %s" % pkg["slug"])
        seen_slugs.add(pkg["slug"])
        packages.append(pkg)

    JSON_OUT.parent.mkdir(parents = True, exist_ok = True)
    JSON_OUT.write_text(json.dumps(packages, indent = 2, ensure_ascii = False), encoding = "utf-8")
    SEED_OUT.write_text(render_seed(packages), encoding = "utf-8")

    print("wrote %s (%s packages)" % (SEED_OUT, len(packages)))
    for pkg in packages:
        print(
            "%s %s price=%s tags=%s"
            % (pkg["sheet"], pkg["slug"], pkg["base_price"], len(pkg["inclustions"]))
        )


if __name__ == "__main__":
    main()
