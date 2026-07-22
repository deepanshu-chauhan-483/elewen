#!/usr/bin/env python3
"""Parse Birthday Packages - Pastel.xlsx into JSON + Strapi seed JS."""

from __future__ import annotations

import json
import re
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "Birthday Packages - Pastel.xlsx"
JSON_OUT = ROOT / "tmp" / "pastel_packages_parsed.json"
SEED_OUT = ROOT / "strapi" / "src" / "seed" / "decor-pastel-packages.js"

# sheet -> (slug, title, theme_name, level, child_theme_slug)
META: dict[str, tuple[str, str, str, str, str]] = {
    "Del1": ("delilah-1", "Delilah 1", "Delilah", "level-1", "delilah"),
    "FF1": ("floral-firsts-1", "Floral Firsts 1", "Floral Firsts", "level-1", "floral-firsts"),
    "FF2": ("floral-firsts-2", "Floral Firsts 2", "Floral Firsts", "level-2", "floral-firsts"),
    "FF3": ("floral-firsts-3", "Floral Firsts 3", "Floral Firsts", "level-3", "floral-firsts"),
    "FF23": ("floral-firsts-2-3", "Floral Firsts 2/3", "Floral Firsts", "level-2", "floral-firsts"),
    "Pop1": ("pastel-pop-1", "Pastel Pop 1", "Pastel Pop", "level-1", "pastel-pop"),
    "ULC": (
        "under-the-lights-canopy-1",
        "Under The Lights Canopy 1",
        "Under The Lights Canopy",
        "level-1",
        "under-the-lights-canopy",
    ),
    "ULC2": (
        "under-the-lights-canopy-2",
        "Under The Lights Canopy 2",
        "Under The Lights Canopy",
        "level-2",
        "under-the-lights-canopy",
    ),
    "PIP1": (
        "pretty-in-pastel-1",
        "Pretty in Pastel 1",
        "Pretty in Pastel",
        "level-1",
        "pretty-in-pastel",
    ),
    "Pip2": (
        "pretty-in-pastel-2",
        "Pretty in Pastel 2",
        "Pretty in Pastel",
        "level-2",
        "pretty-in-pastel",
    ),
    "SS1": ("starlit-spirits-1", "Starlit Spirits 1", "Starlit Spirits", "level-1", "starlit-spirits"),
    "SS2": ("starlit-spirits-2", "Starlit Spirits 2", "Starlit Spirits", "level-2", "starlit-spirits"),
    "WWK": ("words-we-keep-1", "Words We Keep 1", "Words We Keep", "level-1", "words-we-keep"),
    "SAS": ("sip-and-slove-1", "Sip and Slove 1", "Sip and Slove", "level-1", "sip-and-slove"),
    "BB": ("bloom-bar-1", "Bloom Bar 1", "Bloom Bar", "level-1", "bloom-bar"),
}

DEFAULT_NITTY = (
    "Love the look but need a different color? go ahead and book! "
    "We'll contact you to finalize the perfect shade from our available decor to match your vibe."
)
DEFAULT_SETUP = "A small area for our team to work, access to water and electricity."
PRICE_COVERS = [
    {"text": "The Complete Look"},
    {"text": "Professional Setup"},
    {"text": "Final Dismantling"},
    {"text": "Coordination with decor team handled end-to-end to ensure timely setup."},
    {"text": "No hidden service fees."},
]

# Card chips must stay short (Baby Shower style). Excel "What's Included" is often too long.
SHORT_CARD_TAGS: dict[str, list[str]] = {
    "Del1": [
        "5 Small glass vases with bunches of floral arrangements",
        "20 Pillar Pastel or neutral-toned candles with Ribbon detailing on candles",
    ],
    "FF1": ["Custom engraved mirror", "4 big artificial flowers", "Base artificial flowers"],
    "FF2": ["Custom engraved mirror", "Real flower mirror decor", "3 big artificial flowers"],
    "FF3": ["Printed panel backdrop", "Real flower decoration", "Pillar candles in glasses"],
    "FF23": ["Printed panel backdrop", "Real flower decoration", "Pillar candles in glasses"],
    "Pop1": ["Easel stand", "Custom printed board", "Balloon decor on easel", "10 helium balloons"],
    "ULC": [
        "Canopy structure",
        "Flower bunch on canopy",
        "10 fairy light strings",
        "Mattress and carpet",
        "6 cushions",
        "2 lantern cages",
    ],
    "ULC2": [
        "Wooden canopy",
        "5 fairy light strings",
        "Disco balls / lanterns",
        "2 tables",
        "Gadda and carpet",
        "Table fabric draping",
        "8 cushions",
    ],
    "PIP1": ["Single panel backdrop", "Balloon or floral decor", "Happy Birthday neon", "Cake table"],
    "Pip2": ["Double panel backdrop", "Custom message board", "Cake table"],
    "SS1": ["Basic bar", "Custom bar menu + easel", "3 cocktail tables", "3 small centrepieces"],
    "SS2": [
        "Wooden bar",
        "4 poster structure",
        "Draped cloth",
        "Fairy lights",
        "Custom bar menu + easel",
        "3 cocktail tables",
    ],
    "WWK": [
        "Wall panel 8x4ft",
        "Instruction board",
        "Letters",
        "Printed pictures",
        "Marker stand",
        "Markers",
    ],
    "SAS": ["Grass wall", "Custom crossword board", "2 Markers"],
    "BB": [
        "Bloom bar counter",
        "1 stand",
        "8 flowers + 2 greens",
        "Instruction easel",
        "Bouquet paper",
        "Stationery kit",
        "Umbrella",
    ],
}

SECTION_STOPS = [
    "What's Included",
    "Whats Included",
    "Nitty - Gritties",
    "Nitty-Gritties",
    "The price you see covers it all",
    "Best placed",
    "Why we love",
    "Setup",
    "Policies",
    "Cancellation",
    "Reference",
]


def clean_prose(text: str, fallback: str = "") -> str:
    """Drop policy / cancellation / T&C tails that Excel appends into prose cells."""
    cleaned = (text or "").strip()
    if not cleaned:
        return fallback

    lower = cleaned.lower()
    cut_markers = [
        " policies ",
        "policies all the material",
        "cancellation policy",
        "read our full",
        "[read our",
        "reference picture",
    ]
    for marker in cut_markers:
        idx = lower.find(marker)
        if idx > 0:
            cleaned = cleaned[:idx].strip(" .")
            break

    return cleaned[:500] if cleaned else fallback


def shorten_inclusion_tag(text: str) -> str:
    """Keep listing chips short enough to render as pills, not paragraphs."""
    cleaned = re.sub(r"\s+", " ", (text or "").strip(" -"))
    if not cleaned:
        return ""

    lower = cleaned.lower()
    if lower.startswith("reference") or "ideal setup" in lower:
        return ""
    if lower.startswith("policies") or "cancellation" in lower:
        return ""

    # Prefer the head before long parentheticals / dashes.
    cleaned = re.split(r"\(| - | — |; ", cleaned, maxsplit = 1)[0].strip(" /")
    if len(cleaned) > 42:
        cleaned = cleaned[:39].rstrip(" ,/-") + "..."
    return cleaned



def norm(value: object) -> str:
    if value is None:
        return ""
    text = str(value).replace("\r\n", "\n").replace("\r", "\n").strip()
    return re.sub(r"[ \t]+", " ", text)


def amt(value: object) -> float | None:
    if isinstance(value, (int, float)):
        return float(value)
    text = (
        norm(value)
        .replace(",", "")
        .replace("₹", "")
        .replace("Rs.", "")
        .replace("Rs", "")
        .replace("INR", "")
    )
    match = re.search(r"-?\d+(?:\.\d+)?", text)
    return float(match.group(0)) if match else None


def find_label(ws, labels: list[str], start: int = 1) -> int | None:
    targets = {label.lower().rstrip(":") for label in labels}
    for row in range(start, (ws.max_row or 1) + 1):
        cell = norm(ws.cell(row, 1).value).lower().rstrip(":")
        if cell in targets:
            return row
    return None


def amount_on_row(ws, row: int) -> float | None:
    for col in range(2, 8):
        value = amt(ws.cell(row, col).value)
        if value is not None:
            return value
    return None


def section_after(ws, labels: list[str], stop_labels: list[str]) -> str:
    row = find_label(ws, labels)
    if row is None:
        return ""
    stops = {label.lower().rstrip(":") for label in stop_labels}
    lines: list[str] = []
    for next_row in range(row + 1, (ws.max_row or row) + 1):
        text = norm(ws.cell(next_row, 1).value)
        if not text:
            continue
        key = text.lower().rstrip(":")
        if key in stops or any(key.startswith(stop) for stop in stops):
            break
        lines.append(text)
    return "\n".join(lines)


def js_string(value: str) -> str:
    escaped = (
        value.replace("\\", "\\\\")
        .replace("'", "\\'")
        .replace("\n", "\\n")
        .replace("\r", "")
    )
    return "'%s'" % escaped


def parse_sheet(ws, sheet_name: str, meta: tuple[str, str, str, str, str]) -> dict:
    slug, title, theme_name, level, child = meta
    package_price = platform = labour = total = None

    for row in range(1, (ws.max_row or 1) + 1):
        label = norm(ws.cell(row, 1).value).lower().rstrip(":")
        if label == "package price":
            package_price = amount_on_row(ws, row)
        elif "transport" in label and "labour" in label:
            labour = amount_on_row(ws, row)
        elif label == "platform fee":
            platform = amount_on_row(ws, row)
        elif label.startswith("total price"):
            total = amount_on_row(ws, row)

    desc_raw = section_after(
        ws,
        ["Description", "Description "],
        SECTION_STOPS,
    )
    description = clean_prose(
        " ".join(
            line for line in desc_raw.split("\n") if line and not line.startswith("-")
        ).strip(),
        title,
    )

    includes_raw = section_after(
        ws,
        ["What's Included", "Whats Included", "What's Included:"],
        SECTION_STOPS,
    )
    inclusions: list[str] = []
    if sheet_name in SHORT_CARD_TAGS:
        inclusions = list(SHORT_CARD_TAGS[sheet_name])
    else:
        for line in includes_raw.split("\n"):
            text = shorten_inclusion_tag(line.lstrip("- ").strip())
            if text:
                inclusions.append(text)

    if not inclusions:
        for row in range(3, 40):
            label = norm(ws.cell(row, 1).value).lower()
            if label in ("package price", "platform fee") or "transport" in label:
                break
            detail = shorten_inclusion_tag(norm(ws.cell(row, 2).value))
            if detail and detail.lower() not in ("details", "vendor cost", "average"):
                inclusions.append(detail)

    nitty = section_after(
        ws,
        ["Nitty - Gritties", "Nitty-Gritties", "Nitty Gritties"],
        SECTION_STOPS,
    )
    nitty = clean_prose(" ".join(nitty.split("\n")).strip(), DEFAULT_NITTY)

    best = section_after(
        ws,
        ["Best placed at", "Best placed", "Best Placed At"],
        SECTION_STOPS,
    )
    best = clean_prose(" ".join(best.split("\n")).strip(), "")

    why = section_after(
        ws,
        ["Why we love it", "Why we love", "Why We Love It"],
        SECTION_STOPS,
    )
    why = clean_prose(" ".join(why.split("\n")).strip(), "")

    setup = section_after(
        ws,
        ["Setup needs", "Setup", "Set up needs"],
        SECTION_STOPS,
    )
    setup = clean_prose(
        " ".join(
            line for line in setup.split("\n") if not line.lower().startswith("the price")
        ).strip(),
        DEFAULT_SETUP,
    )

    breakdown: list[dict[str, int]] = []
    for row in range(3, 50):
        label = norm(ws.cell(row, 1).value).lower().rstrip(":")
        if (
            label in ("package price", "platform fee")
            or label.startswith("total price")
            or ("transport" in label and "labour" in label)
        ):
            break
        detail = norm(ws.cell(row, 2).value)
        cost = amt(ws.cell(row, 3).value)
        if (
            detail
            and cost is not None
            and cost >= 100
            and detail.lower() not in ("details", "vendor cost", "average")
        ):
            breakdown.append({"label": detail[:80], "amount": int(round(cost))})

    if platform is not None:
        breakdown.append({"label": "Platform fee", "amount": int(round(platform))})
    if labour is not None:
        breakdown.append({"label": "Labour & Transport", "amount": int(round(labour))})

    if total is not None:
        base_price = int(round(total))
    else:
        base_price = int(round((package_price or 0) + (platform or 0) + (labour or 0)))

    key_decor = clean_prose(description, title)[:180]

    return {
        "sheet": sheet_name,
        "title": title,
        "slug": slug,
        "theme_name": theme_name,
        "budget_level_slug": level,
        "base_price": base_price,
        "description": description[:1000],
        "key_decor_elements": key_decor,
        "nitty_gritty": nitty[:800],
        "setup_needs": setup[:500],
        "best_placed_at": best[:500],
        "why_we_love_it": why[:500],
        "child_theme": child,
        "price_breakdown": breakdown,
        "inclustions": inclusions[:12],
    }


def render_seed(packages: list[dict]) -> str:
    chunks: list[str] = [
        "'use strict';",
        "",
        "const {",
        "  ensureBudgetLevelDocumentId,",
        "  ensureEventDocumentId,",
        "  ensureThemeDocumentId,",
        "  seedDecorPackages,",
        "} = require('./decor-proposal-shared');",
        "",
        "const EVENT_SLUG_CANDIDATES = [",
        "  'birthday-parties',",
        "  'birthday-celebrations',",
        "  'birthday',",
        "  'birthday-party',",
        "];",
        "",
        "const PASTEL_PACKAGES = [",
    ]

    for pkg in packages:
        lines = ["  {"]
        lines.append("    sheet: %s," % js_string(pkg["sheet"]))
        lines.append("    title: %s," % js_string(pkg["title"]))
        lines.append("    slug: %s," % js_string(pkg["slug"]))
        lines.append("    theme_name: %s," % js_string(pkg["theme_name"]))
        lines.append("    budget_level_slug: %s," % js_string(pkg["budget_level_slug"]))
        lines.append("    base_price: %s," % pkg["base_price"])
        lines.append("    description: %s," % js_string(pkg["description"]))
        lines.append("    key_decor_elements: %s," % js_string(pkg["key_decor_elements"]))
        lines.append("    nitty_gritty: %s," % js_string(pkg["nitty_gritty"]))
        lines.append("    setup_needs: %s," % js_string(pkg["setup_needs"]))
        lines.append("    best_placed_at: %s," % js_string(pkg["best_placed_at"]))
        lines.append("    why_we_love_it: %s," % js_string(pkg["why_we_love_it"]))
        lines.append(
            "    theme_slug_candidates: ['pastel', %s],"
            % js_string(pkg["child_theme"])
        )
        lines.append("    price_covers: [")
        for cover in PRICE_COVERS:
            lines.append("      { text: %s }," % js_string(cover["text"]))
        lines.append("    ],")
        lines.append("    price_breakdown: [")
        for item in pkg["price_breakdown"]:
            lines.append(
                "      { label: %s, amount: %s },"
                % (js_string(item["label"]), item["amount"])
            )
        lines.append("    ],")
        lines.append("    inclustions: [")
        for text in pkg["inclustions"]:
            lines.append(
                "      { text: %s, group_title: 'Includes' }," % js_string(text)
            )
        lines.append("    ],")
        lines.append("  },")
        chunks.append("\n".join(lines))

    chunks.extend(
        [
            "];",
            "",
            "async function seedDecorPastelPackages(strapi) {",
            "  await ensureThemeDocumentId(strapi, { name: 'Pastel', slug: 'pastel', display_order: 60 });",
            "  await ensureEventDocumentId(strapi, {",
            "    name: 'Birthday Parties',",
            "    slugCandidates: EVENT_SLUG_CANDIDATES,",
            "    display_order: 10,",
            "  });",
            "  await ensureBudgetLevelDocumentId(strapi, {",
            "    name: 'Level 1',",
            "    slug: 'level-1',",
            "    display_order: 1,",
            "    description: 'Thoughtfully put together essentials for smaller guest counts.',",
            "  });",
            "  await ensureBudgetLevelDocumentId(strapi, {",
            "    name: 'Level 2',",
            "    slug: 'level-2',",
            "    display_order: 2,",
            "    description: 'A more detailed setup with a stronger visual and service experience.',",
            "  });",
            "  await ensureBudgetLevelDocumentId(strapi, {",
            "    name: 'Level 3',",
            "    slug: 'level-3',",
            "    display_order: 3,",
            "    description: 'Premium styling and service depth for larger or more layered celebrations.',",
            "  });",
            "",
            "  return seedDecorPackages(strapi, PASTEL_PACKAGES, {",
            "    eventSlugCandidates: EVENT_SLUG_CANDIDATES,",
            "    logPrefix: 'seed/decor-pastel',",
            "    refreshExistingDecor: true,",
            "  });",
            "}",
            "",
            "module.exports = {",
            "  PASTEL_PACKAGES,",
            "  EVENT_SLUG_CANDIDATES,",
            "  seedDecorPastelPackages,",
            "};",
            "",
        ]
    )
    return "\n".join(chunks)


def main() -> None:
    workbook = load_workbook(XLSX, data_only = True)
    packages: list[dict] = []

    for sheet_name, meta in META.items():
        if sheet_name not in workbook.sheetnames:
            raise SystemExit("Missing sheet: %s" % sheet_name)
        packages.append(parse_sheet(workbook[sheet_name], sheet_name, meta))

    JSON_OUT.parent.mkdir(parents = True, exist_ok = True)
    JSON_OUT.write_text(json.dumps(packages, indent = 2, ensure_ascii = False), encoding = "utf-8")
    SEED_OUT.write_text(render_seed(packages), encoding = "utf-8")

    print("wrote %s (%s packages)" % (SEED_OUT, len(packages)))
    for pkg in packages:
        print(
            "%s %s price=%s inclusions=%s"
            % (pkg["sheet"], pkg["slug"], pkg["base_price"], len(pkg["inclustions"]))
        )


if __name__ == "__main__":
    main()
