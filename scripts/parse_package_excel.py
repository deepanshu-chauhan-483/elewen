#!/usr/bin/env python3
"""Parse Elewen package Excel workbooks into structured JSON summaries."""

from __future__ import annotations

import json
import re
import sys
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.worksheet.worksheet import Worksheet

ROOT = Path(__file__).resolve().parents[1]
SEED_DIR = ROOT / "strapi" / "src" / "seed"
OUTPUT_PATH = ROOT / "tmp" / "excel_package_summary.json"

WORKBOOKS: list[dict[str, Any]] = [
    {
        "file": "Birthday _ Anniversary - Bubbles & Black Tie.xlsx",
        "exclude_patterns": [],
    },
    {
        "file": "Engagement - Sufi Night Package.xlsx",
        "exclude_patterns": [],
    },
    {
        "file": "General Pack _ House Party _ Birthday Party.xlsx",
        "exclude_patterns": [r"^pincodes", r"^copy"],
    },
    {
        "file": "Final Catering Menu - 14.05.26.xlsx",
        "exclude_patterns": [],
        "structure_only": True,
    },
]


def normalize_text(value: object) -> str:
    """Return trimmed string for a cell value."""
    if value is None:
        return ""
    text = str(value).replace("\r\n", "\n").replace("\r", "\n").strip()
    return re.sub(r"\s+", " ", text)


def slugify(text: str) -> str:
    """Build a URL-friendly slug from title text."""
    slug = text.lower()
    slug = re.sub(r"[^a-z0-9]+", "-", slug)
    slug = re.sub(r"-+", "-", slug).strip("-")
    return slug or "untitled"


def parse_amount(value: object) -> float | None:
    """Parse numeric amount from spreadsheet cell."""
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = normalize_text(value)
    text = text.replace(",", "").replace("₹", "").replace("Rs.", "").replace("Rs", "")
    text = text.replace("INR", "").strip()
    match = re.search(r"-?\d+(?:\.\d+)?", text)
    if not match:
        return None
    return float(match.group(0))


def row_texts(ws: Worksheet, row_idx: int, max_col: int | None = None) -> list[str]:
    """Collect non-empty cell texts for one row."""
    limit = max_col or ws.max_column or 1
    texts: list[str] = []
    for col_idx in range(1, limit + 1):
        text = normalize_text(ws.cell(row = row_idx, column = col_idx).value)
        if text:
            texts.append(text)
    return texts


def find_label_row(
    ws: Worksheet,
    label: str,
    start_row: int = 1,
    end_row: int | None = None,
) -> int | None:
    """Find first row whose first non-empty cell matches label (case-insensitive)."""
    last_row = end_row or ws.max_row or 1
    target = label.lower().strip()
    for row_idx in range(start_row, last_row + 1):
        texts = row_texts(ws, row_idx, max_col = 3)
        if not texts:
            continue
        if texts[0].lower().rstrip(":") == target:
            return row_idx
    return None


def find_amount_after_label(
    ws: Worksheet,
    labels: list[str],
    start_row: int = 1,
    end_row: int | None = None,
) -> float | None:
    """Find amount in same row as one of the labels."""
    last_row = end_row or ws.max_row or 1
    normalized_labels = {label.lower().rstrip(":") for label in labels}
    for row_idx in range(start_row, last_row + 1):
        texts = row_texts(ws, row_idx, max_col = 6)
        if not texts:
            continue
        if texts[0].lower().rstrip(":") not in normalized_labels:
            continue
        for text in texts[1:]:
            amount = parse_amount(text)
            if amount is not None:
                return amount
    return None


def collect_paragraph_after(
    ws: Worksheet,
    start_row: int,
    stop_labels: set[str],
    max_rows: int = 20,
) -> str:
    """Collect first long paragraph starting after start_row."""
    chunks: list[str] = []
    for offset in range(1, max_rows + 1):
        row_idx = start_row + offset
        if row_idx > (ws.max_row or 0):
            break
        texts = row_texts(ws, row_idx, max_col = 4)
        if not texts:
            if chunks:
                break
            continue
        first = texts[0]
        if first.lower().rstrip(":") in stop_labels:
            break
        chunks.append(first if len(texts) == 1 else " ".join(texts))
        if len(" ".join(chunks)) > 120:
            break
    return " ".join(chunks).strip()


def collect_section_text(
    ws: Worksheet,
    section_label: str,
    stop_labels: set[str],
    max_rows: int = 30,
    join_rows: str = " ",
) -> str:
    """Collect multi-line section body until next known label."""
    row_idx = find_label_row(ws, section_label)
    if row_idx is None:
        return ""
    chunks: list[str] = []
    for offset in range(1, max_rows + 1):
        next_row = row_idx + offset
        if next_row > (ws.max_row or 0):
            break
        texts = row_texts(ws, next_row, max_col = 6)
        if not texts:
            if chunks:
                break
            continue
        first = texts[0]
        if first.lower().rstrip(":") in stop_labels:
            break
        chunks.append(" ".join(texts))
    if join_rows == "\n\n":
        return "\n\n".join(chunks).strip()
    return " ".join(chunks).strip()


def detect_estimated_cost_column(ws: Worksheet) -> int:
    """Return column index for Estimated cost header, defaulting to column C."""
    for row_idx in range(1, 6):
        for col_idx in range(1, 8):
            header = normalize_text(ws.cell(row = row_idx, column = col_idx).value).lower()
            if header in {"estimated cost", "estimated cost average"}:
                return col_idx
    return 3


def read_priced_amount(ws: Worksheet, row_idx: int, cost_col: int) -> float | None:
    """Read line-item cost from the estimated-cost column, with tier-column fallback."""
    primary = parse_amount(ws.cell(row = row_idx, column = cost_col).value)
    if primary is not None and primary >= 100:
        return primary
    for col_idx in range(3, 7):
        if col_idx == cost_col:
            continue
        candidate = parse_amount(ws.cell(row = row_idx, column = col_idx).value)
        if candidate is not None and candidate >= 100:
            return candidate
    return None


def extract_line_items(ws: Worksheet, start_row: int, end_row: int) -> list[dict[str, Any]]:
    """Extract priced line items before totals block."""
    items: list[dict[str, Any]] = []
    cost_col = detect_estimated_cost_column(ws)
    skip_labels = {
        "total",
        "platform fee",
        "transport",
        "transport fee",
        "labour",
        "labor",
        "labour & transport",
        "total price",
        "package price",
        "description",
        "includes",
        "includes:",
        "add-ons",
        "add ons",
        "addons",
        "level 1",
        "level 2",
        "level 3",
    }
    for row_idx in range(start_row, end_row + 1):
        col_a = normalize_text(ws.cell(row = row_idx, column = 1).value)
        col_b = normalize_text(ws.cell(row = row_idx, column = 2).value)
        label = col_a or col_b
        if not label:
            continue
        label_key = label.lower().rstrip(":")
        if label_key in skip_labels:
            continue
        amount = read_priced_amount(ws, row_idx, cost_col)
        details = col_b if col_a and col_b else ""
        if amount is None and not details:
            continue
        items.append(
            {
                "label": label,
                "details": details,
                "amount": amount,
            }
        )
    return items


def extract_tagline_description(ws: Worksheet) -> str:
    """Extract marketing tagline paragraph after the pricing block."""
    anchor_row = find_label_row(ws, "Total Price") or find_label_row(ws, "Total price")
    if anchor_row is None:
        return ""
    stop_tokens = ("key decor", "includes:", "includes", "the nitty-gritty")
    for offset in range(2, 10):
        row_idx = anchor_row + offset
        if row_idx > (ws.max_row or 0):
            break
        text = normalize_text(ws.cell(row = row_idx, column = 1).value)
        if not text or len(text) < 24:
            continue
        lowered = text.lower()
        if any(token in lowered for token in stop_tokens):
            break
        if not text.endswith(":"):
            return text
    return ""


def extract_key_decor_elements(ws: Worksheet) -> str:
    """Extract Key Decor Elements body text."""
    row_idx = find_label_row(ws, "Key Decor Elements")
    if row_idx is None:
        return ""
    chunks: list[str] = []
    for offset in range(1, 5):
        next_row = row_idx + offset
        if next_row > (ws.max_row or 0):
            break
        text = normalize_text(ws.cell(row = next_row, column = 1).value)
        if not text:
            break
        lowered = text.lower().rstrip(":")
        if lowered in {"includes", "includes:"}:
            break
        chunks.append(text)
    return " ".join(chunks).strip()


def extract_elewen_placement_blurb(ws: Worksheet) -> str:
    """Extract Elewen placement recommendation used as why-we-love-it copy."""
    last_row = ws.max_row or 1
    for row_idx in range(1, last_row + 1):
        text = normalize_text(ws.cell(row = row_idx, column = 1).value)
        if not text:
            continue
        lowered = text.lower()
        if lowered.startswith("elewen thinks") or lowered.startswith("this is the") or lowered.startswith("this is for"):
            return text
    return ""


def extract_addons(ws: Worksheet) -> list[dict[str, Any]]:
    """Extract add-on rows from sheet."""
    addons: list[dict[str, Any]] = []
    addon_row = find_label_row(ws, "Add-ons") or find_label_row(ws, "Add ons")
    if addon_row is None:
        addon_row = find_label_row(ws, "Addons")
    if addon_row is None:
        return addons
    stop_labels = {
        "description",
        "includes",
        "includes:",
        "the nitty-gritty",
        "nitty-gritty",
        "setup",
        "best placed",
        "why we love it",
    }
    for offset in range(1, 40):
        row_idx = addon_row + offset
        if row_idx > (ws.max_row or 0):
            break
        texts = row_texts(ws, row_idx, max_col = 6)
        if not texts:
            if addons:
                break
            continue
        if texts[0].lower().rstrip(":") in stop_labels:
            break
        amount = None
        details = ""
        for text in texts[1:]:
            parsed = parse_amount(text)
            if parsed is not None and amount is None:
                amount = parsed
            elif not details:
                details = text
        addons.append({"label": texts[0], "details": details, "amount": amount})
    return addons


def extract_choice_hints(includes_text: str) -> list[str]:
    """Find choose/select hints inside includes text."""
    hints: list[str] = []
    patterns = [
        r"choose\s+[^.\n]+",
        r"select\s+[^.\n]+",
        r"pick\s+[^.\n]+",
        r"\(\s*choose[^)]+\)",
        r"\d+\s*\+\s*\d+",
    ]
    for pattern in patterns:
        for match in re.finditer(pattern, includes_text, flags = re.IGNORECASE):
            hint = normalize_text(match.group(0))
            if hint and hint not in hints:
                hints.append(hint)
    return hints


def should_exclude_sheet(sheet_name: str, exclude_patterns: list[str]) -> bool:
    """Return True when sheet should be skipped."""
    lowered = sheet_name.strip().lower()
    for pattern in exclude_patterns:
        if re.search(pattern, lowered, flags = re.IGNORECASE):
            return True
    reference_names = {"cp", "copy", "pincodes", "pincode", "reference", "menu index"}
    if lowered in reference_names:
        return True
    if lowered.startswith("copy"):
        return True
    if lowered.startswith("pincodes"):
        return True
    return False


def parse_package_sheet(ws: Worksheet, sheet_name: str) -> dict[str, Any]:
    """Parse one package worksheet into structured summary."""
    title = ""
    for row_idx in range(1, min(15, (ws.max_row or 1) + 1)):
        texts = row_texts(ws, row_idx, max_col = 3)
        if not texts:
            continue
        if texts[0].lower() in {"details", "estimated cost", "average"}:
            continue
        if len(texts[0]) > 2 and not texts[0].lower().startswith("total"):
            title = texts[0]
            break

    total_row = find_label_row(ws, "Total", start_row = 1, end_row = 40)
    total_price_row = find_label_row(ws, "Total Price", start_row = 1, end_row = 60)
    line_item_end = (total_row or total_price_row or 20) - 1

    package_price = find_amount_after_label(ws, ["Package Price"], start_row = 1, end_row = 50)
    if package_price is None:
        package_price = find_amount_after_label(ws, ["Total"], start_row = 1, end_row = 40)
    platform_fee = find_amount_after_label(ws, ["Platform fee", "Platform Fee"], start_row = 1, end_row = 50)
    transport = find_amount_after_label(
        ws,
        ["Transport", "Transport Fee", "Labour", "Labor", "Labour & Transport", "Transport / Labour"],
        start_row = 1,
        end_row = 50,
    )
    total_price = find_amount_after_label(ws, ["Total Price", "Total price"], start_row = 1, end_row = 60)

    tagline = extract_tagline_description(ws)
    key_decor_elements = extract_key_decor_elements(ws)
    placement_blurb = extract_elewen_placement_blurb(ws)

    stop_labels = {
        "includes",
        "includes:",
        "add-ons",
        "add ons",
        "addons",
        "the nitty-gritty",
        "nitty-gritty",
        "setup",
        "best placed",
        "why we love it",
        "the price you see covers it all",
    }
    description = ""
    for label in ["Description", "Decription"]:
        description_row = find_label_row(ws, label)
        if description_row is not None:
            description = collect_paragraph_after(ws, description_row, stop_labels)
            if description:
                break
    if not description and tagline:
        description = tagline

    includes_text = collect_section_text(ws, "Includes", stop_labels, join_rows = "\n\n")
    if not includes_text:
        includes_text = collect_section_text(ws, "Includes:", stop_labels, join_rows = "\n\n")

    nitty_gritty = collect_section_text(ws, "The Nitty-Gritty", stop_labels)
    if not nitty_gritty:
        nitty_gritty = collect_section_text(ws, "Nitty-Gritty", stop_labels)

    setup = collect_section_text(ws, "Setup Needs", stop_labels)
    if not setup:
        setup = collect_section_text(ws, "Setup", stop_labels)
    best_placed = collect_section_text(ws, "Best Placed at", stop_labels)
    if not best_placed:
        best_placed = collect_section_text(ws, "Best placed", stop_labels)
    why_section = collect_section_text(ws, "Why we love it", stop_labels)
    why_we_love_it = placement_blurb or why_section

    line_items = extract_line_items(ws, start_row = 2, end_row = max(line_item_end, 2))
    addons = extract_addons(ws)
    choice_hints = extract_choice_hints(includes_text)

    suggested_slug = slugify(title or sheet_name)

    return {
        "sheet_name": sheet_name,
        "title": title or sheet_name,
        "suggested_slug": suggested_slug,
        "package_price": package_price,
        "platform_fee": platform_fee,
        "transport_labour": transport,
        "total_price": total_price,
        "missing_total_price": total_price is None,
        "missing_package_price": package_price is None,
        "description": description or tagline,
        "tagline": tagline,
        "key_decor_elements": key_decor_elements,
        "includes": includes_text,
        "nitty_gritty": nitty_gritty,
        "setup": setup,
        "best_placed": best_placed,
        "why_we_love_it": why_we_love_it,
        "line_items": line_items,
        "addons": addons,
        "choice_group_hints": choice_hints,
    }


def analyze_catering_workbook(ws_names: list[str], wb: openpyxl.Workbook) -> dict[str, Any]:
    """Summarize catering menu workbook structure."""
    sheet_summaries: list[dict[str, Any]] = []
    looks_like_menu = False
    looks_like_packages = False

    for sheet_name in ws_names:
        ws = wb[sheet_name]
        headers: list[str] = []
        sample_rows: list[list[str]] = []
        for row_idx in range(1, min(6, (ws.max_row or 1) + 1)):
            texts = row_texts(ws, row_idx, max_col = 8)
            if row_idx == 1:
                headers = texts
            elif texts:
                sample_rows.append(texts)

        header_blob = " ".join(headers).lower()
        if any(token in header_blob for token in ["dish", "course", "category", "menu", "veg", "non veg"]):
            looks_like_menu = True
        if any(token in header_blob for token in ["package", "price", "total price", "platform fee"]):
            looks_like_packages = True

        sheet_summaries.append(
            {
                "sheet_name": sheet_name,
                "headers": headers,
                "sample_rows": sample_rows[:3],
                "rows": ws.max_row,
                "cols": ws.max_column,
            }
        )

    existing_catering_slugs = load_existing_slugs(["catering-*.js"])
    comparison = {
        "workbook_type": (
            "menu_library"
            if looks_like_menu and not looks_like_packages
            else "mixed"
            if looks_like_menu and looks_like_packages
            else "unknown"
        ),
        "looks_like_menu_data": looks_like_menu,
        "looks_like_package_rows": looks_like_packages,
        "existing_catering_seed_files": sorted(existing_catering_slugs.keys()),
        "note": (
            "Compare dish/category sheets to strapi/src/seed/catering-menu-data.js "
            "and package-like rows to catering-essentials.js / catering-assisted.js / catering-end-to-end.js."
        ),
    }

    return {
        "structure_analysis": comparison,
        "sheets": sheet_summaries,
        "packages": [],
    }


def load_existing_slugs(glob_patterns: list[str]) -> dict[str, list[str]]:
    """Load slug values from seed JS files."""
    slug_map: dict[str, list[str]] = {}
    for pattern in glob_patterns:
        for path in SEED_DIR.glob(pattern):
            content = path.read_text(encoding = "utf-8")
            slugs = re.findall(r"slug:\s*['\"]([^'\"]+)['\"]", content)
            package_slugs = [
                slug
                for slug in slugs
                if slug not in {"decor", "catering"}
                and not slug.startswith("level-")
            ]
            if package_slugs:
                slug_map[path.name] = sorted(set(package_slugs))
    return slug_map


def find_duplicate_slugs(parsed_slugs: list[str]) -> dict[str, list[str]]:
    """Map parsed slugs to existing seed files with same slug."""
    all_existing: dict[str, list[str]] = {}
    for path in SEED_DIR.glob("*.js"):
        content = path.read_text(encoding = "utf-8")
        for slug in re.findall(r"slug:\s*['\"]([^'\"]+)['\"]", content):
            if slug in {"decor", "catering"} or slug.startswith("level-"):
                continue
            all_existing.setdefault(slug, []).append(path.name)

    duplicates: dict[str, list[str]] = {}
    for slug in parsed_slugs:
        if slug in all_existing:
            duplicates[slug] = sorted(set(all_existing[slug]))
    return duplicates


def parse_workbook(spec: dict[str, Any]) -> dict[str, Any]:
    """Parse one workbook spec."""
    file_path = ROOT / spec["file"]
    wb = openpyxl.load_workbook(file_path, data_only = True)
    exclude_patterns: list[str] = spec.get("exclude_patterns", [])

    if spec.get("structure_only"):
        return {
            "file": spec["file"],
            **analyze_catering_workbook(wb.sheetnames, wb),
        }

    packages: list[dict[str, Any]] = []
    for sheet_name in wb.sheetnames:
        if should_exclude_sheet(sheet_name, exclude_patterns):
            continue
        ws = wb[sheet_name]
        parsed = parse_package_sheet(ws, sheet_name)
        packages.append(parsed)

    slugs = [pkg["suggested_slug"] for pkg in packages]
    return {
        "file": spec["file"],
        "packages": packages,
        "duplicate_slug_matches": find_duplicate_slugs(slugs),
    }


def main() -> int:
    """Run parser and write JSON summary."""
    OUTPUT_PATH.parent.mkdir(parents = True, exist_ok = True)
    results: list[dict[str, Any]] = []
    for spec in WORKBOOKS:
        file_path = ROOT / spec["file"]
        if not file_path.exists():
            print("skip missing workbook:", spec["file"])
            continue
        results.append(parse_workbook(spec))

    payload = {
        "generated_at": "2026-07-17",
        "output_path": str(OUTPUT_PATH),
        "workbooks": results,
        "all_existing_seed_slugs": load_existing_slugs(["*.js"]),
    }

    OUTPUT_PATH.write_text(json.dumps(payload, indent = 2, ensure_ascii = False), encoding = "utf-8")
    print(str(OUTPUT_PATH))
    return 0


if __name__ == "__main__":
    sys.exit(main())
